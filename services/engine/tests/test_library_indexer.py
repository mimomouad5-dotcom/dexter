# Component: n29 مستورد وفهرسة المكتبة & n16 مدير المهام (Unit Tests)
"""
Tests for Phase 6:
- Magic bytes classification (PPTX, PDF, PNG, Text, Corrupt, Unsupported)
- 3-stage deduplication pipeline
- Template metadata extraction (slide count, aspect ratio, shapes, titles)
- Cooperative cancellation safety and transactional rollback
"""

import os
import tempfile
import zipfile
import pytest
from pptx import Presentation
from pptx.util import Inches

from dexter_engine.storage.db import DatabaseManager
from dexter_engine.tasks.task_manager import TaskManager, TaskState
from dexter_engine.library.magic_classifier import classify_file
from dexter_engine.library.deduplicator import DeduplicationPipeline, calculate_full_sha256
from dexter_engine.library.template_extractor import extract_presentation_metadata
from dexter_engine.library.indexer import LibraryIndexer

@pytest.fixture
def temp_workspace():
    with tempfile.TemporaryDirectory() as tmpdir:
        db_path = os.path.join(tmpdir, "test.db")
        db = DatabaseManager(db_path)
        tasks = TaskManager(db)
        indexer = LibraryIndexer(db, tasks)
        yield {
            "root": tmpdir,
            "db": db,
            "tasks": tasks,
            "indexer": indexer
        }

def test_magic_classifier_file_types(temp_workspace):
    root = temp_workspace["root"]

    # 1. Valid PNG
    png_path = os.path.join(root, "sample.png")
    with open(png_path, "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR...")
    res = classify_file(png_path)
    assert res["type"] == "image"
    assert res["subtype"] == "png"
    assert res["status"] == "valid"

    # 2. Valid PDF
    pdf_path = os.path.join(root, "sample.pdf")
    with open(pdf_path, "wb") as f:
        f.write(b"%PDF-1.4\n1 0 obj\n<<>>\nendobj")
    res = classify_file(pdf_path)
    assert res["type"] == "document"
    assert res["subtype"] == "pdf"
    assert res["status"] == "valid"

    # 3. Text / Markdown
    md_path = os.path.join(root, "notes.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("# Presentation Notes\nSome bullet points here.")
    res = classify_file(md_path)
    assert res["type"] == "text"
    assert res["subtype"] == "markdown"
    assert res["status"] == "valid"

    # 4. Corrupt Zip masquerading as PPTX
    corrupt_path = os.path.join(root, "broken.pptx")
    with open(corrupt_path, "wb") as f:
        f.write(b"PK\x03\x04corrupted_payload_without_central_directory")
    res = classify_file(corrupt_path)
    assert res["type"] == "corrupt"
    assert res["status"] == "corrupt"

    # 5. Unsupported binary executable
    bin_path = os.path.join(root, "malicious.exe")
    with open(bin_path, "wb") as f:
        f.write(b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00")
    res = classify_file(bin_path)
    assert res["type"] == "unsupported"
    assert res["status"] == "unsupported"

def test_template_metadata_extraction(temp_workspace):
    root = temp_workspace["root"]
    pptx_path = os.path.join(root, "test_template.pptx")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # 16:9
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "مقدمة المعمارية السحابية"
    prs.save(pptx_path)

    meta = extract_presentation_metadata(pptx_path)
    assert meta["slide_count"] == 1
    assert meta["aspect_ratio"] == "16:9"
    assert "مقدمة المعمارية السحابية" in meta["titles"]
    assert meta["is_template_candidate"] is True
    assert meta["error"] is None

def test_deduplication_pipeline(temp_workspace):
    root = temp_workspace["root"]
    file1 = os.path.join(root, "file1.txt")
    file2 = os.path.join(root, "file2_dup.txt")

    content = b"Exact identical presentation slide content repeated twice."
    with open(file1, "wb") as f:
        f.write(content)
    with open(file2, "wb") as f:
        f.write(content)

    pipe = DeduplicationPipeline()
    sha1, is_dup1, orig1 = pipe.process_file_hash(file1, len(content))
    sha2, is_dup2, orig2 = pipe.process_file_hash(file2, len(content))

    assert is_dup1 is False
    assert is_dup2 is True
    assert orig2 == file1
    assert sha1 == sha2

def test_full_library_scan_and_query(temp_workspace):
    root = temp_workspace["root"]
    indexer: LibraryIndexer = temp_workspace["indexer"]
    scan_dir = os.path.join(root, "user_library")
    os.makedirs(scan_dir, exist_ok=True)

    # Populate directory with files
    prs = Presentation()
    prs.save(os.path.join(scan_dir, "template_a.pptx"))
    
    # Duplicate of template_a
    with open(os.path.join(scan_dir, "template_a.pptx"), "rb") as src:
        data = src.read()
    with open(os.path.join(scan_dir, "template_a_copy.pptx"), "wb") as dst:
        dst.write(data)

    # Add a note
    with open(os.path.join(scan_dir, "readme.md"), "w", encoding="utf-8") as f:
        f.write("# Readme Notes")

    # Run scan
    res = indexer.scan_directory(scan_dir)
    assert res["status"] == "completed"
    assert res["processed"] == 3
    assert res["duplicates"] == 1
    assert res["valid"] == 2

    # Query files
    all_files = indexer.query_files()
    assert len(all_files) == 3

    # Query duplicates
    dup_clusters = indexer.get_duplicates()
    assert len(dup_clusters) == 1
    assert dup_clusters[0]["count"] == 2

def test_scan_cooperative_cancellation(temp_workspace):
    root = temp_workspace["root"]
    indexer: LibraryIndexer = temp_workspace["indexer"]
    tasks: TaskManager = temp_workspace["tasks"]

    scan_dir = os.path.join(root, "cancel_test_dir")
    os.makedirs(scan_dir, exist_ok=True)

    # Create 5 files
    for i in range(5):
        with open(os.path.join(scan_dir, f"doc_{i}.txt"), "w") as f:
            f.write(f"Sample content {i}")

    task_id = tasks.create_task("library_scan")
    # Request cancellation before scan starts
    tasks.request_cancellation(task_id)

    res = indexer.scan_directory(scan_dir, task_id=task_id)
    assert res["status"] == "cancelled"
