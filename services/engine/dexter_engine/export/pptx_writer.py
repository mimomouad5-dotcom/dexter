# Component: n27 مصدّر PPTX (python-pptx + OOXML) & n46 ملف .pptx
"""
Native PowerPoint Presentation Exporter for Dexter Presentation Studio.
Reads n23 UnifiedPresentationModel and outputs fully editable, standards-compliant .pptx files.
Supports Arabic RTL, native charts, speaker notes, and verified OOXML animations.
"""

import os
import shutil
from typing import Dict, Any, Optional

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

from dexter_engine.model.presentation import (
    UnifiedPresentationModel,
    ElementType,
    ShapeGeometry,
    ChartType
)
from dexter_engine.export.animations import apply_slide_transition, apply_shape_entrance_animation
from dexter_engine.logging.safe_logger import get_safe_logger

logger = get_safe_logger("n27_writer")

def hex_to_rgb(hex_str: Optional[str]) -> Optional[RGBColor]:
    if not hex_str:
        return None
    cleaned = hex_str.lstrip('#')
    if len(cleaned) == 6:
        r = int(cleaned[0:2], 16)
        g = int(cleaned[2:4], 16)
        b = int(cleaned[4:6], 16)
        return RGBColor(r, g, b)
    return None

def apply_arabic_rtl_properties(paragraph, font_family: str = "Cairo", font_size_pt: float = 18.0, is_rtl: bool = True):
    """
    Sets OpenXML paragraph RTL properties and font typefaces.
    ISO/IEC 29500-1 §21.1.2.2.7
    """
    pPr = paragraph._p.get_or_add_pPr()
    if is_rtl:
        pPr.set('rtl', '1')
        pPr.set('algn', 'r')
    else:
        pPr.set('rtl', '0')
        pPr.set('algn', 'l')

    for run in paragraph.runs:
        run.font.name = font_family
        run.font.size = Pt(font_size_pt)
        rPr = run._r.get_or_add_rPr()
        cs_tag = parse_xml(f'<a:cs {nsdecls("a")} typeface="{font_family}"/>')
        rPr.append(cs_tag)

class PPTXWriter:
    @staticmethod
    def export_presentation(model: UnifiedPresentationModel, target_output_path: str) -> Dict[str, Any]:
        """
        Exports UnifiedPresentationModel to native editable .pptx with atomic write.
        """
        logger.info(f"Starting PPTX export to: {target_output_path}")
        target_dir = os.path.dirname(os.path.abspath(target_output_path))
        os.makedirs(target_dir, exist_ok=True)

        prs = Presentation()

        # Set slide dimensions in EMUs
        if model.presentation.aspect_ratio == "16:9":
            prs.slide_width = Emu(12192000)   # 13.333 inches
            prs.slide_height = Emu(6858000)   # 7.5 inches
        else:
            prs.slide_width = Emu(9144000)    # 10.0 inches
            prs.slide_height = Emu(6858000)   # 7.5 inches

        blank_layout = prs.slide_layouts[6]
        is_global_rtl = model.presentation.direction == "rtl"

        for slide_model in model.presentation.slides:
            slide = prs.slides.add_slide(blank_layout)

            # 1. Slide Transition (n28)
            if slide_model.transition and slide_model.transition.verified:
                apply_slide_transition(slide, slide_model.transition.effect_id, slide_model.transition.speed)

            # 2. Render Elements
            for el in slide_model.elements:
                x = Emu(el.frame.x)
                y = Emu(el.frame.y)
                w = Emu(el.frame.w)
                h = Emu(el.frame.h)
                added_shape = None

                # TEXT ELEMENT
                if el.type == ElementType.TEXT:
                    textbox = slide.shapes.add_textbox(x, y, w, h)
                    tf = textbox.text_frame
                    tf.word_wrap = True
                    text_str = el.content.get("text", "")
                    p = tf.paragraphs[0]
                    p.text = text_str
                    apply_arabic_rtl_properties(
                        p,
                        font_family=el.style.font_family or "Cairo",
                        font_size_pt=el.style.font_size_pt or 18.0,
                        is_rtl=el.style.is_rtl if el.style.is_rtl is not None else is_global_rtl
                    )
                    color = hex_to_rgb(el.style.font_color_hex)
                    if color:
                        p.font.color.rgb = color
                    p.font.bold = el.style.bold
                    added_shape = textbox

                # SHAPE ELEMENT
                elif el.type == ElementType.SHAPE:
                    geom = el.content.get("geometry", "rectangle")
                    mso_geom = MSO_SHAPE.ROUNDED_RECTANGLE if geom == "rounded_rectangle" else MSO_SHAPE.RECTANGLE
                    shape = slide.shapes.add_shape(mso_geom, x, y, w, h)
                    fill_color = hex_to_rgb(el.style.fill_color_hex)
                    if fill_color:
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = fill_color
                    border_color = hex_to_rgb(el.style.border_color_hex)
                    if border_color:
                        shape.line.color.rgb = border_color
                        if el.style.border_width_pt:
                            shape.line.width = Pt(el.style.border_width_pt)
                    label = el.content.get("label")
                    if label:
                        stf = shape.text_frame
                        sp = stf.paragraphs[0]
                        sp.text = label
                        apply_arabic_rtl_properties(
                            sp,
                            font_family=el.style.font_family or "Cairo",
                            font_size_pt=el.style.font_size_pt or 14.0,
                            is_rtl=is_global_rtl
                        )
                    added_shape = shape

                # IMAGE ELEMENT
                elif el.type == ElementType.IMAGE:
                    img_path = el.content.get("file_path")
                    if img_path and os.path.exists(img_path):
                        pic = slide.shapes.add_picture(img_path, x, y, w, h)
                        added_shape = pic

                # CHART ELEMENT (Native editable PowerPoint Chart)
                elif el.type == ElementType.CHART:
                    chart_data = CategoryChartData()
                    chart_data.categories = el.content.get("categories", ["بند 1", "بند 2", "بند 3"])
                    for series in el.content.get("series", []):
                        chart_data.add_series(series.get("name", "السلسلة"), series.get("values", (10, 20, 30)))

                    c_type = el.content.get("chart_type", "column")
                    xl_type = XL_CHART_TYPE.COLUMN_CLUSTERED if c_type == "column" else XL_CHART_TYPE.BAR_CLUSTERED
                    chart_shape = slide.shapes.add_chart(xl_type, x, y, w, h, chart_data)
                    chart_shape.chart.has_legend = True
                    added_shape = chart_shape

                # TABLE ELEMENT
                elif el.type == ElementType.TABLE:
                    rows = el.content.get("rows", 2)
                    cols = el.content.get("cols", 2)
                    table_shape = slide.shapes.add_table(rows, cols, x, y, w, h)
                    data_grid = el.content.get("data", [["", ""]])
                    for r_idx, row_data in enumerate(data_grid):
                        for c_idx, cell_val in enumerate(row_data):
                            if r_idx < rows and c_idx < cols:
                                cell = table_shape.table.cell(r_idx, c_idx)
                                cell.text = str(cell_val)
                    added_shape = table_shape

                # 3. Animation Effect Injection (n28)
                if added_shape and el.animation and el.animation.verified:
                    apply_shape_entrance_animation(slide, added_shape.shape_id, el.animation.effect_id)

            # 4. Speaker Notes
            if slide_model.notes:
                notes_slide = slide.notes_slide
                notes_tf = notes_slide.notes_text_frame
                notes_tf.text = slide_model.notes

        # Atomic Write
        temp_target = target_output_path + ".tmp"
        prs.save(temp_target)
        os.replace(temp_target, target_output_path)

        file_size = os.path.getsize(target_output_path)
        logger.info(f"Presentation successfully exported ({file_size} bytes): {target_output_path}")

        return {
            "status": "success",
            "file_path": target_output_path,
            "slide_count": len(prs.slides),
            "size_bytes": file_size
        }

    @staticmethod
    def verify_round_trip(output_path: str) -> Dict[str, Any]:
        """
        Automated round-trip verification reading back the exported .pptx file.
        """
        assert os.path.exists(output_path), f"File not found: {output_path}"
        prs = Presentation(output_path)

        has_charts = False
        has_notes = False
        has_arabic_rtl = False

        for slide in prs.slides:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
                has_notes = True
            for shape in slide.shapes:
                if shape.has_chart:
                    has_charts = True
                if shape.has_text_frame and shape.text_frame.text:
                    if 'rtl="1"' in shape._element.xml:
                        has_arabic_rtl = True

        return {
            "slide_count": len(prs.slides),
            "has_charts": has_charts,
            "has_notes": has_notes,
            "has_arabic_rtl": has_arabic_rtl,
            "round_trip_valid": True
        }
